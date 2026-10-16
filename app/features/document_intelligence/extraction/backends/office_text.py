"""Fast, search-grade text from office formats: Word, Excel and PowerPoint
(OOXML and 97-2003), OpenDocument, RTF and saved email.

The words and figures, one paragraph block per line, read with pure-Python
parsers in milliseconds; layout, styles and table structure are not kept.
It is deliberately not in the default registry, where the document pipeline,
Docling and XlsxBackend own these formats with full structure. A caller that
needs speed over structure (Google Drive search reranking twenty files
within seconds) registers it explicitly.

The legacy binary formats are read straight from their record structures
(olefile), so nothing shells out to an office suite.
"""

from __future__ import annotations

import email
import importlib
import io
import re
import struct
import time
import zipfile
from datetime import date, datetime
from email import policy
from typing import BinaryIO, Callable
from xml.etree import ElementTree

from app.features.document_intelligence.canonical.enums import (
    BlockType,
    ExtractionErrorCode,
    ExtractionMethod,
    ExtractionStatus,
)
from app.features.document_intelligence.canonical.schemas import (
    CanonicalDocument,
    DocumentBlock,
    ExtractionMetadata,
    Provenance,
    SourceFile,
)
from app.features.document_intelligence.extraction.interfaces import (
    ExtractionBackend,
)
from app.features.document_intelligence.extraction.limits import (
    ResourceLimitExceeded,
    check_paragraph_count,
    check_zip_container,
)
from app.features.document_intelligence.extraction.schemas import (
    ExtractionCapabilities,
    ExtractionError,
    ExtractionOptions,
    ExtractionResult,
)


class UnreadableOfficeFile(Exception):
    """The file cannot be read: encrypted, corrupt, or not what it claims."""


# The parsers the readers import on first use: about a second together
# (openpyxl alone half of it), otherwise paid inside the first request that
# meets each format.
_PARSERS = ("docx", "openpyxl", "pptx", "xlrd", "olefile", "striprtf.striprtf")


def preload() -> None:
    """Import every reader's parser now, before serving. One that is missing
    is skipped: its reader fails its own format on use, as before."""
    for name in _PARSERS:
        try:
            importlib.import_module(name)
        except ImportError:
            pass


# Rows read per sheet. A sheet past this is a data dump, and its first rows
# carry the headers and the shape that search needs.
_MAX_ROWS = 5000


class OfficeTextBackend(ExtractionBackend):

    _ZIP_EXTENSIONS = {".docx", ".xlsx", ".xlsm", ".pptx", ".odt", ".ods", ".odp"}

    @property
    def name(self) -> str:
        return "office_text"

    @property
    def capabilities(self) -> ExtractionCapabilities:
        return ExtractionCapabilities(
            supported_mime_types={
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                "application/msword",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "application/vnd.ms-excel.sheet.macroEnabled.12",
                "application/vnd.ms-excel",
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                "application/vnd.ms-powerpoint",
                "application/vnd.oasis.opendocument.text",
                "application/vnd.oasis.opendocument.spreadsheet",
                "application/vnd.oasis.opendocument.presentation",
                "application/rtf",
                "text/rtf",
                "message/rfc822",
            },
            supports_ocr=False,
            supports_tables=False,
            supports_images=False,
            supports_layout=False,
            supports_provenance=True,
        )

    def supports(self, source: SourceFile) -> bool:
        return source.extension.lower() in _READERS

    async def extract(
        self,
        source: SourceFile,
        content: BinaryIO,
        options: ExtractionOptions,
    ) -> ExtractionResult:
        started_at = time.perf_counter()
        extension = source.extension.lower()

        try:
            content.seek(0)
            data = content.read()

            if not data:
                return self._failed(ExtractionErrorCode.INVALID_CONTENT, "The file is empty.")

            if extension in self._ZIP_EXTENSIONS:
                check_zip_container(io.BytesIO(data))

            text = _READERS[extension](data)
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            check_paragraph_count(len(lines))

            document = CanonicalDocument(
                id=f"document:{source.content_hash or source.file_id}",
                source=source,
                blocks=[
                    DocumentBlock(
                        id=f"office:paragraph:{ordinal}",
                        type=BlockType.PARAGRAPH,
                        text=line,
                        ordinal=ordinal,
                        provenance=[
                            Provenance(
                                source_element_id=f"line:{ordinal}",
                                extraction_method=ExtractionMethod.NATIVE,
                            )
                        ],
                    )
                    for ordinal, line in enumerate(lines)
                ],
                metadata={"source_backend": self.name},
                extraction=ExtractionMetadata(
                    backend=self.name,
                    status=ExtractionStatus.SUCCESS,
                    duration_ms=int((time.perf_counter() - started_at) * 1000),
                ),
            )

            return ExtractionResult(status=ExtractionStatus.SUCCESS, document=document)

        except ResourceLimitExceeded as exc:
            return self._failed(ExtractionErrorCode.RESOURCE_LIMIT_EXCEEDED, str(exc))
        except UnreadableOfficeFile as exc:
            return self._failed(ExtractionErrorCode.PARSING_FAILED, str(exc))
        except Exception as exc:
            return self._failed(
                ExtractionErrorCode.PARSING_FAILED,
                f"Office text extraction failed for '{source.filename}'.",
                {"exception_type": type(exc).__name__, "exception": str(exc)},
            )

    def _failed(self, code: ExtractionErrorCode, message: str, details: dict | None = None) -> ExtractionResult:
        return ExtractionResult(
            status=ExtractionStatus.FAILED,
            errors=[
                ExtractionError(
                    code=code,
                    message=message,
                    retryable=False,
                    backend=self.name,
                    details=details or {},
                )
            ],
        )


# ---------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------


def _docx(content: bytes) -> str:
    """Paragraphs, then table rows: tables carry much of what people search
    for (figures, dates, owners)."""
    from docx import Document

    document = Document(io.BytesIO(content))

    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            line = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if line:
                parts.append(line)

    return "\n".join(parts)


def _doc(content: bytes) -> str:
    """The text of a Word 97-2003 file, from its piece table.

    A .doc stores its text in runs ("pieces"), each 8-bit or UTF-16, listed
    in the CLX structure of the table stream. Reading the pieces in order
    recovers the text exactly, where scanning for printable runs picks up
    style names and garbage.
    """
    import olefile

    with olefile.OleFileIO(io.BytesIO(content)) as ole:
        word = ole.openstream("WordDocument").read()
        flags = struct.unpack_from("<H", word, 0x0A)[0]

        if flags & 0x0100:
            raise UnreadableOfficeFile("encrypted .doc")

        table = ole.openstream("1Table" if flags & 0x0200 else "0Table").read()

    # Walk the FIB's variable-length arrays to fcClx/lcbClx, pair 33 of the
    # fc/lcb array.
    pos = 32
    csw = struct.unpack_from("<H", word, pos)[0]
    pos += 2 + csw * 2
    cslw = struct.unpack_from("<H", word, pos)[0]
    pos += 2 + cslw * 4
    pos += 2  # cbRgFcLcb
    fc_clx, lcb_clx = struct.unpack_from("<II", word, pos + 33 * 8)

    pos, end = fc_clx, fc_clx + lcb_clx
    plc = b""

    while pos < end:
        kind = table[pos]

        if kind == 0x01:  # Prc: formatting, skipped
            pos += 3 + struct.unpack_from("<H", table, pos + 1)[0]
        elif kind == 0x02:  # Pcdt: the piece table
            size = struct.unpack_from("<I", table, pos + 1)[0]
            plc = table[pos + 5 : pos + 5 + size]
            break
        else:
            raise UnreadableOfficeFile("unrecognised .doc piece table")

    pieces = (len(plc) - 4) // 12

    if pieces <= 0:
        raise UnreadableOfficeFile(".doc has no text pieces")

    cps = struct.unpack_from(f"<{pieces + 1}I", plc, 0)
    parts: list[str] = []

    for i in range(pieces):
        raw = struct.unpack_from("<I", plc, 4 * (pieces + 1) + 8 * i + 2)[0]
        count = cps[i + 1] - cps[i]

        if raw & 0x40000000:
            start = (raw & 0x3FFFFFFF) // 2
            parts.append(word[start : start + count].decode("cp1252", errors="replace"))
        else:
            start = raw & 0x3FFFFFFF
            parts.append(word[start : start + 2 * count].decode("utf-16-le", errors="replace"))

    text = "".join(parts)
    # Fields are \x13 instructions \x14 result \x15: keep the result.
    text = re.sub(r"\x13[^\x13\x14\x15]*\x14", "", text)
    text = re.sub(r"\x13[^\x13\x14\x15]*\x15", "", text)
    text = text.replace("\x07", " | ")  # table cell marks

    return _paragraphs(text.split("\r"))


# ---------------------------------------------------------------------
# Spreadsheets
# ---------------------------------------------------------------------


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat() if not value.time() else value.isoformat(" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _rows_text(title: str, rows) -> list[str]:
    lines = [f"Sheet: {title}"]

    for count, row in enumerate(rows):
        if count >= _MAX_ROWS:
            break

        # Empty cells stay, so each value is under its own column's header;
        # only the trailing ones go.
        cells = [_cell(v) for v in row]

        while cells and not cells[-1]:
            cells.pop()

        if cells:
            lines.append(" | ".join(cells))

    return lines


def _xlsx(content: bytes) -> str:
    from openpyxl import load_workbook

    # data_only: a formula's last computed value, which is what the person
    # saw and might search for.
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)

    try:
        lines: list[str] = []

        for sheet in book.worksheets:
            lines += _rows_text(sheet.title, sheet.iter_rows(values_only=True))

        return "\n".join(lines)
    finally:
        book.close()


def _xls(content: bytes) -> str:
    import xlrd

    book = xlrd.open_workbook(file_contents=content, on_demand=True)
    lines: list[str] = []

    for sheet in book.sheets():
        def values(sheet=sheet):
            for r in range(sheet.nrows):
                row = []

                for cell in sheet.row(r):
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                            continue
                        except Exception:
                            pass
                    row.append(cell.value)

                yield row

        lines += _rows_text(sheet.name, values())

    return "\n".join(lines)


# ---------------------------------------------------------------------
# Presentations
# ---------------------------------------------------------------------


def _pptx(content: bytes) -> str:
    from pptx import Presentation

    deck = Presentation(io.BytesIO(content))
    lines: list[str] = []

    def shape_text(shape):
        if shape.has_text_frame and shape.text_frame.text.strip():
            lines.append(shape.text_frame.text.strip())

        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                line = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if line:
                    lines.append(line)

        for child in getattr(shape, "shapes", ()):  # groups
            shape_text(child)

    for number, slide in enumerate(deck.slides, 1):
        lines.append(f"Slide {number}")

        for shape in slide.shapes:
            shape_text(shape)

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Notes: {notes}")

    return "\n".join(lines)


# PowerPoint 97-2003 records: an 8-byte header (version and instance, type,
# length); a version of 0xF marks a container whose body is more records.
_PPT_TEXT_CHARS = 0x0FA0   # UTF-16
_PPT_TEXT_BYTES = 0x0FA8   # 8-bit
# Masters and handouts hold placeholder text that is in every deck.
_PPT_SKIP = {0x03F8, 0x0FC9}


def _ppt(content: bytes) -> str:
    import olefile

    with olefile.OleFileIO(io.BytesIO(content)) as ole:
        data = ole.openstream("PowerPoint Document").read()

    texts: list[str] = []
    stack = [(0, len(data))]

    while stack:
        pos, end = stack.pop()

        while pos + 8 <= end:
            ver_inst, rec_type, length = struct.unpack_from("<HHI", data, pos)
            body, pos = pos + 8, pos + 8 + length

            if (ver_inst & 0x000F) == 0x000F:
                if rec_type not in _PPT_SKIP:
                    # Resumed after this container, so text keeps document order.
                    stack.append((pos, end))
                    stack.append((body, min(pos, end)))
                    break
            elif rec_type == _PPT_TEXT_CHARS:
                texts.append(data[body:pos].decode("utf-16-le", errors="replace"))
            elif rec_type == _PPT_TEXT_BYTES:
                texts.append(data[body:pos].decode("cp1252", errors="replace"))

    return _paragraphs(texts)


# ---------------------------------------------------------------------
# OpenDocument
# ---------------------------------------------------------------------

_ODF_TEXT = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
_ODF_TABLE = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"


def _odf_inline(element) -> str:
    """An element's text, with ODF's encoded spaces and breaks restored."""
    out = [element.text or ""]

    for child in element:
        tag = child.tag.rsplit("}", 1)[-1]

        if tag == "s":
            out.append(" " * int(child.get(f"{{{_ODF_TEXT}}}c", "1")))
        elif tag in ("tab", "line-break"):
            out.append(" ")
        else:
            out.append(_odf_inline(child))

        out.append(child.tail or "")

    return "".join(out)


def _odf(content: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        root = ElementTree.fromstring(archive.read("content.xml"))

    lines: list[str] = []
    paragraph_tags = {f"{{{_ODF_TEXT}}}p", f"{{{_ODF_TEXT}}}h"}
    row_tag = f"{{{_ODF_TABLE}}}table-row"
    cell_tags = {f"{{{_ODF_TABLE}}}table-cell", f"{{{_ODF_TABLE}}}covered-table-cell"}

    def walk(element):
        for child in element:
            if child.tag == row_tag:
                cells = []

                for cell in child:
                    if cell.tag in cell_tags:
                        text = " ".join(
                            _odf_inline(p).strip() for p in cell.iter() if p.tag in paragraph_tags
                        ).strip()
                        if text:
                            cells.append(text)

                if cells:
                    lines.append(" | ".join(cells))
            elif child.tag in paragraph_tags:
                text = _odf_inline(child).strip()
                if text:
                    lines.append(text)
            else:
                walk(child)

    walk(root)

    return "\n".join(lines)


# ---------------------------------------------------------------------
# RTF and email
# ---------------------------------------------------------------------


def _rtf(content: bytes) -> str:
    from striprtf.striprtf import rtf_to_text

    # RTF is 7-bit with escapes for everything else; latin-1 maps every
    # byte, so a stray 8-bit byte cannot fail the decode.
    return rtf_to_text(content.decode("latin-1"), errors="ignore").strip()


def _eml(content: bytes) -> str:
    message = email.message_from_bytes(content, policy=policy.default)
    lines = [
        f"{header}: {message[header]}"
        for header in ("From", "To", "Cc", "Date", "Subject")
        if message[header]
    ]

    body = message.get_body(preferencelist=("plain", "html"))

    if body is not None:
        text = body.get_content()

        if body.get_content_type() == "text/html":
            text = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", text, flags=re.S | re.I)
            text = re.sub(r"<[^>]+>", " ", text)

        lines.append(text.strip())

    attachments = [name for name in (part.get_filename() for part in message.iter_attachments()) if name]

    if attachments:
        lines.append("Attachments: " + ", ".join(attachments))

    return "\n".join(lines)


def _paragraphs(chunks) -> str:
    """Clean paragraphs, each once: legacy formats repeat text freely."""
    seen: set[str] = set()
    out: list[str] = []

    for chunk in chunks:
        for line in re.split(r"[\r\n\x0b\x0c]+", chunk):
            line = re.sub(r"[\x00-\x08\x0e-\x1f]", "", line).strip()

            if line and line not in seen:
                seen.add(line)
                out.append(line)

    return "\n".join(out)


_READERS: dict[str, Callable[[bytes], str]] = {
    ".docx": _docx,
    ".doc": _doc,
    ".xlsx": _xlsx,
    ".xlsm": _xlsx,
    ".xls": _xls,
    ".pptx": _pptx,
    ".ppt": _ppt,
    ".odt": _odf,
    ".ods": _odf,
    ".odp": _odf,
    ".rtf": _rtf,
    ".eml": _eml,
}
