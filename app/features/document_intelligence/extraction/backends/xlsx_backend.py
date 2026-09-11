from __future__ import annotations

import time
import zipfile
from io import BytesIO
from typing import BinaryIO

import openpyxl

from app.features.document_intelligence.canonical.enums import (
    BlockType,
    ExtractionErrorCode,
    ExtractionMethod,
    ExtractionStatus,
)
from app.features.document_intelligence.canonical.schemas import (
    CanonicalDocument,
    DocumentBlock,
    ExtractionMetadata,
    Provenance,
    SourceFile,
    TableCell,
    TableData,
)
from app.features.document_intelligence.extraction.interfaces import (
    ExtractionBackend,
)
from app.features.document_intelligence.extraction.limits import (
    ResourceLimitExceeded,
    check_table_dimensions,
    check_worksheet_count,
    check_zip_container,
)
from app.features.document_intelligence.extraction.schemas import (
    ExtractionCapabilities,
    ExtractionError,
    ExtractionOptions,
    ExtractionResult,
)


class XlsxBackend(ExtractionBackend):
    """Native extraction backend for XLSX/XLSM workbooks.

    Uses openpyxl in normal (not read_only) mode. read_only mode was
    tested and confirmed to drop merged-cell information entirely
    (ReadOnlyWorksheet has no merged_cells attribute at all), which
    would sacrifice the exact fidelity advantage that justifies using
    openpyxl over pandas.read_excel in the first place. At the
    platform's 25MB upload cap, normal mode's memory footprint is a
    non-issue. .xlsm is the identical OOXML zip structure with a
    macro-enabled content type, so it's handled by the same code path
    with no reason to reconsider read_only mode just for this format.

    Docling also supports .xlsx, but was tested and found to silently
    drop sheet names (it maps sheets to an internal page number with no
    heading/caption anywhere in the output), so a multi-sheet workbook
    becomes anonymous tables with no way to tell which sheet they came
    from. openpyxl gives direct access to sheet names.

    VBA macro content is never executed and is not read: openpyxl's
    default load (keep_vba=False) discards the vbaProject.bin part
    entirely, so macros are effectively opaque and untouched by this
    backend.

    Since .xlsx/.xlsm are zip containers, the input size cap alone
    doesn't protect against a zip-bomb-style file; that's checked
    before opening the workbook.
    """

    _SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}

    @property
    def name(self) -> str:
        return "xlsx"

    @property
    def capabilities(self) -> ExtractionCapabilities:
        return ExtractionCapabilities(
            supported_mime_types={
                (
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                "application/vnd.ms-excel.sheet.macroEnabled.12",
            },
            supports_ocr=False,
            supports_tables=True,
            supports_images=False,
            supports_layout=False,
            supports_provenance=True,
        )

    def supports(self, source: SourceFile) -> bool:
        return source.extension.lower() in self._SUPPORTED_EXTENSIONS

    async def extract(
        self,
        source: SourceFile,
        content: BinaryIO,
        options: ExtractionOptions,
    ) -> ExtractionResult:
        started_at = time.perf_counter()
        workbook = None

        try:
            content.seek(0)
            data = content.read()

            if not data:
                return ExtractionResult(
                    status=ExtractionStatus.FAILED,
                    errors=[
                        ExtractionError(
                            code=ExtractionErrorCode.INVALID_CONTENT,
                            message="The file is empty.",
                            retryable=False,
                            backend=self.name,
                        )
                    ],
                )

            try:
                check_zip_container(BytesIO(data))
            except zipfile.BadZipFile as exc:
                return ExtractionResult(
                    status=ExtractionStatus.FAILED,
                    errors=[
                        ExtractionError(
                            code=ExtractionErrorCode.INVALID_CONTENT,
                            message=f"Not a valid XLSX file: {exc}",
                            retryable=False,
                            backend=self.name,
                        )
                    ],
                )

            workbook = openpyxl.load_workbook(
                BytesIO(data),
                read_only=False,
                data_only=True,
            )

            check_worksheet_count(len(workbook.sheetnames))

            blocks = self._build_blocks(workbook)

            document = CanonicalDocument(
                id=self._document_id(source),
                source=source,
                blocks=blocks,
                metadata={"source_backend": self.name},
                extraction=ExtractionMetadata(
                    backend=self.name,
                    status=ExtractionStatus.SUCCESS,
                    duration_ms=int(
                        (time.perf_counter() - started_at) * 1000
                    ),
                    used_ocr=False,
                    used_vision=False,
                ),
            )

            return ExtractionResult(
                status=ExtractionStatus.SUCCESS,
                document=document,
            )

        except ResourceLimitExceeded as exc:
            return ExtractionResult(
                status=ExtractionStatus.FAILED,
                errors=[
                    ExtractionError(
                        code=ExtractionErrorCode.RESOURCE_LIMIT_EXCEEDED,
                        message=str(exc),
                        retryable=False,
                        backend=self.name,
                    )
                ],
            )
        except Exception as exc:
            return ExtractionResult(
                status=ExtractionStatus.FAILED,
                errors=[
                    ExtractionError(
                        code=ExtractionErrorCode.PARSING_FAILED,
                        message=(
                            f"XLSX extraction failed for "
                            f"'{source.filename}'."
                        ),
                        retryable=False,
                        backend=self.name,
                        details={
                            "exception_type": type(exc).__name__,
                            "exception": str(exc),
                        },
                    )
                ],
            )
        finally:
            if workbook is not None:
                workbook.close()

    def _build_blocks(
        self,
        workbook: openpyxl.Workbook,
    ) -> list[DocumentBlock]:
        blocks: list[DocumentBlock] = []
        ordinal = 0

        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]

            blocks.append(
                DocumentBlock(
                    id=f"xlsx:sheet:{sheet_name}:heading",
                    type=BlockType.HEADING,
                    text=sheet_name,
                    heading_level=1,
                    ordinal=ordinal,
                    provenance=[
                        Provenance(
                            source_element_id=f"sheet:{sheet_name}",
                            extraction_method=ExtractionMethod.NATIVE,
                        )
                    ],
                )
            )
            ordinal += 1

            table = self._build_table(worksheet, sheet_name)

            if table is not None:
                blocks.append(
                    DocumentBlock(
                        id=f"xlsx:sheet:{sheet_name}:table",
                        type=BlockType.TABLE,
                        table=table,
                        ordinal=ordinal,
                        provenance=[
                            Provenance(
                                source_element_id=(
                                    f"sheet:{sheet_name}:table"
                                ),
                                extraction_method=(
                                    ExtractionMethod.NATIVE
                                ),
                            )
                        ],
                    )
                )
                ordinal += 1

        return blocks

    def _build_table(
        self,
        worksheet,
        sheet_name: str,
    ) -> TableData | None:
        max_row = worksheet.max_row or 0
        max_col = worksheet.max_column or 0

        if max_row == 0 or max_col == 0:
            return None

        check_table_dimensions(
            rows=max_row,
            columns=max_col,
            context=f"XLSX sheet '{sheet_name}'",
        )

        grid: list[list[object]] = [
            [None] * max_col
            for _ in range(max_row)
        ]

        for r_idx, row in enumerate(
            worksheet.iter_rows(
                min_row=1,
                max_row=max_row,
                max_col=max_col,
            )
        ):
            for c_idx, cell in enumerate(row):
                grid[r_idx][c_idx] = cell.value

        if all(value is None for row in grid for value in row):
            return None

        span_for_top_left: dict[tuple[int, int], tuple[int, int]] = {}
        covered: set[tuple[int, int]] = set()

        for merged_range in worksheet.merged_cells.ranges:
            top_left = (merged_range.min_row, merged_range.min_col)

            row_span = (
                merged_range.max_row - merged_range.min_row + 1
            )
            column_span = (
                merged_range.max_col - merged_range.min_col + 1
            )

            span_for_top_left[top_left] = (row_span, column_span)

            for row_num in range(
                merged_range.min_row,
                merged_range.max_row + 1,
            ):
                for col_num in range(
                    merged_range.min_col,
                    merged_range.max_col + 1,
                ):
                    if (row_num, col_num) != top_left:
                        covered.add((row_num, col_num))

        # Treat the first row as the header, unless it's a single cell
        # merged across the full column width (a title banner rather
        # than column labels) — skip past any such leading rows.
        header_row_index = 0
        for r_idx in range(max_row):
            span = span_for_top_left.get((r_idx + 1, 1))
            if span is not None and span[1] == max_col:
                continue
            header_row_index = r_idx
            break

        cells: list[TableCell] = []

        for r_idx in range(max_row):
            for c_idx in range(max_col):
                row_num = r_idx + 1
                col_num = c_idx + 1

                if (row_num, col_num) in covered:
                    continue

                value = grid[r_idx][c_idx]
                row_span, column_span = span_for_top_left.get(
                    (row_num, col_num),
                    (1, 1),
                )

                cells.append(
                    TableCell(
                        row=r_idx,
                        column=c_idx,
                        text="" if value is None else str(value),
                        row_span=row_span,
                        column_span=column_span,
                        is_header=(r_idx == header_row_index),
                        provenance=[
                            Provenance(
                                source_element_id=(
                                    f"{sheet_name}!R{row_num}C{col_num}"
                                ),
                                extraction_method=(
                                    ExtractionMethod.NATIVE
                                ),
                            )
                        ],
                    )
                )

        return TableData(
            rows=max_row,
            columns=max_col,
            cells=cells,
        )

    def _document_id(self, source: SourceFile) -> str:
        if source.content_hash:
            return f"document:{source.content_hash}"

        return f"document:{source.file_id}"
